import importlib.util
import tempfile
import unittest
from pathlib import Path

SPEC = importlib.util.spec_from_file_location("worker", Path(__file__).parents[1] / "python/worker/worker.py")
worker = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(worker)


class DatasetParsingTests(unittest.TestCase):
    def test_csv_utf8_bom_and_quotes(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "data.csv"
            path.write_text('\ufefftext,source\n"Merhaba, dünya",test\n', encoding="utf-8")
            self.assertEqual(worker.read_rows(path), [{"text": "Merhaba, dünya", "source": "test"}])

    def test_xlsx_header_and_rows(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "data.xlsx"
            book = Workbook(); sheet = book.active
            sheet.append(["text", "source"]); sheet.append(["Deneme", "dosya"]); book.save(path)
            self.assertEqual(worker.read_rows(path), [{"text": "Deneme", "source": "dosya"}])

    def test_extract_txt_docx_and_xlsx(self):
        from openpyxl import Workbook
        from docx import Document
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            text = root / "policy.txt"
            text.write_text("İade süresi 14 gündür.", encoding="utf-8")
            self.assertIn("14 gündür", worker.extract_training_text(text))
            docx = root / "policy.docx"
            doc = Document(); doc.add_paragraph("Kanun maddesi 5."); doc.save(docx)
            self.assertIn("Kanun maddesi 5", worker.extract_training_text(docx))
            xlsx = root / "policy.xlsx"
            book = Workbook(); book.active.append(["Madde", "İçerik"]); book.active.append(["5", "İzin verilir"]); book.save(xlsx)
            self.assertIn("İzin verilir", worker.extract_training_text(xlsx))

    def test_non_rag_does_not_claim_training(self):
        with self.assertRaisesRegex(RuntimeError, "Eğitim işi için model ve veri dosyası seçilmelidir"):
            worker.run({"type": "training"})


if __name__ == "__main__":
    unittest.main()
